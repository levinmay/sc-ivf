"""Antalya source preparation and baseline metadata."""

from __future__ import annotations

from hashlib import sha256
import numpy as np
import pandas as pd
from pathlib import Path
import hashlib
import json
import math
from sklearn.model_selection import GroupShuffleSplit

CATEGORICAL = ["primary_infertility", "diagnosis", "stimulation_type", "trigger_type"]


def numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def prepare_sheet(frame: pd.DataFrame, sheet: str) -> pd.DataFrame:
    rename = {
        "pno": "source_cycle_id",
        "tarih": "retrieval_date",
        "age": "age",
        "inf yrs": "inf_years",
        "prime inf": "primary_infertility",
        "prev ET numb": "previous_et_number",
        "afc": "afc",
        "bmi": "bmi",
        "diag": "diagnosis",
        "stim type": "stimulation_type",
        "trigger type": "trigger_type",
        "stim days": "stimulation_days",
        "FSH": "daily_fsh",
        "hmg": "daily_hmg",
        "tot FSH": "total_fsh",
        "eggs": "eggs",
        "ıcsı": "mii",
        "mat rate": "reported_maturation_rate",
        "FET#": "fet_number",
    }
    out = frame.rename(columns=rename).copy()
    if "reported_maturation_rate" not in out:
        out["reported_maturation_rate"] = np.nan
    if "fet_number" not in out:
        out["fet_number"] = 0
    out["source_sheet"] = sheet
    out["source_row"] = np.arange(len(out), dtype=int)
    keep = [
        "source_cycle_id",
        "retrieval_date",
        "source_sheet",
        "source_row",
        "fet_number",
        "eggs",
        "mii",
        "reported_maturation_rate",
        "age",
        "inf_years",
        "primary_infertility",
        "previous_et_number",
        "afc",
        "bmi",
        "diagnosis",
        "stimulation_type",
        "trigger_type",
        "stimulation_days",
        "daily_fsh",
        "daily_hmg",
        "total_fsh",
    ]
    return out[keep]


def load_antalya(workbook) -> tuple[pd.DataFrame, dict]:
    main_raw = pd.read_excel(workbook, sheet_name="tot cycles LB outcome", header=1)
    no_fet_raw = pd.read_excel(workbook, sheet_name="No FET plus unknown", header=0)
    main_first = main_raw.loc[numeric(main_raw["FET#"]).eq(1)].copy()
    main = prepare_sheet(main_first, "tot cycles LB outcome: FET#=1")
    no_fet = prepare_sheet(no_fet_raw, "No FET plus unknown")
    combined = pd.concat([main, no_fet], ignore_index=True)
    numeric_columns = [
        "source_cycle_id",
        "fet_number",
        "eggs",
        "mii",
        "reported_maturation_rate",
        "age",
        "inf_years",
        "previous_et_number",
        "afc",
        "bmi",
        "stimulation_days",
        "daily_fsh",
        "daily_hmg",
        "total_fsh",
    ]
    for column in numeric_columns:
        combined[column] = numeric(combined[column])
    combined["retrieval_date"] = pd.to_datetime(
        combined["retrieval_date"], errors="coerce"
    )
    for column in CATEGORICAL:
        combined[column] = combined[column].map(
            lambda value: (
                np.nan
                if pd.isna(value) or not str(value).strip()
                else str(value).strip().casefold()
            )
        )
    valid_integer = (
        combined["eggs"].notna()
        & combined["mii"].notna()
        & np.isclose(combined["eggs"], np.round(combined["eggs"]))
        & np.isclose(combined["mii"], np.round(combined["mii"]))
    )
    valid_range = combined["eggs"].gt(0) & combined["mii"].between(0, combined["eggs"])
    valid = valid_integer & valid_range
    cohort = combined.loc[valid].copy().reset_index(drop=True)
    cohort["eggs"] = cohort["eggs"].astype(int)
    cohort["mii"] = cohort["mii"].astype(int)
    cohort["log1p_eggs"] = np.log1p(cohort["eggs"])
    if cohort["source_cycle_id"].isna().any() or cohort["retrieval_date"].isna().any():
        raise RuntimeError("Antalya source ID or retrieval date is missing")
    cohort["source_couple_id"] = np.floor(cohort["source_cycle_id"]).astype(int)
    cohort["Patient_Code"] = cohort["source_couple_id"].map(
        lambda value: f"ANT-{value:05d}"
    )
    cohort["retrieval_id"] = (
        cohort["Patient_Code"] + "-" + cohort["retrieval_date"].dt.strftime("%Y%m%d")
    )
    date_ordinal = cohort["retrieval_date"].map(pd.Timestamp.toordinal).astype(float)
    cohort["retrieval_date_linear"] = (date_ordinal - date_ordinal.mean()) / 365.25
    day_of_year = cohort["retrieval_date"].dt.dayofyear.to_numpy(float)
    cohort["retrieval_season_sin"] = np.sin(2 * np.pi * day_of_year / 365.25)
    cohort["retrieval_season_cos"] = np.cos(2 * np.pi * day_of_year / 365.25)
    cohort["egg_stratum"] = pd.cut(
        cohort["eggs"],
        bins=[0, 5, 10, 15, 20, np.inf],
        labels=["1-5", "6-10", "11-15", "16-20", "21+"],
        right=True,
    ).astype(str)
    reported = combined["reported_maturation_rate"].notna() & valid
    rate_diff = (
        combined.loc[reported, "reported_maturation_rate"]
        - combined.loc[reported, "mii"] / combined.loc[reported, "eggs"]
    ).abs()
    digest = sha256(workbook.read_bytes()).hexdigest()
    couple_counts = cohort.groupby("Patient_Code").size()
    duplicate_signature_columns = [
        "Patient_Code",
        "eggs",
        "mii",
        "age",
        "inf_years",
        "previous_et_number",
        "afc",
        "bmi",
        "diagnosis",
        "stimulation_type",
        "trigger_type",
        "stimulation_days",
        "daily_fsh",
        "daily_hmg",
        "total_fsh",
    ]
    duplicated_signature = cohort.duplicated(duplicate_signature_columns, keep=False)
    implausible = {
        "bmi_outside_12_60": int(
            (~cohort.bmi.between(12, 60) & cohort.bmi.notna()).sum()
        ),
        "afc_outside_0_80": int(
            (~cohort.afc.between(0, 80) & cohort.afc.notna()).sum()
        ),
        "stimulation_days_outside_5_20": int(
            (
                ~cohort.stimulation_days.between(5, 20)
                & cohort.stimulation_days.notna()
            ).sum()
        ),
    }
    audit = {
        "workbook_sha256": digest,
        "main_all_fet_rows": int(len(main_raw)),
        "main_first_fet_rows": int(len(main_first)),
        "main_later_fet_rows_excluded": int(len(main_raw) - len(main_first)),
        "no_fet_or_unknown_rows": int(len(no_fet_raw)),
        "released_retrieval_records_before_validity_filter": int(len(combined)),
        "valid_external_cycles": int(len(cohort)),
        "unique_retrieval_ids": int(cohort.retrieval_id.nunique()),
        "apparent_unique_couples": int(cohort.Patient_Code.nunique()),
        "apparent_couples_with_two_retrievals": int(couple_counts.gt(1).sum()),
        "maximum_retrievals_per_apparent_couple": int(couple_counts.max()),
        "rows_in_possible_duplicate_signatures": int(duplicated_signature.sum()),
        "invalid_count_rows": int((~valid).sum()),
        "rows_with_reported_maturation_rate": int(reported.sum()),
        "max_abs_reported_rate_discrepancy": float(rate_diff.max())
        if len(rate_diff)
        else None,
        "published_total_retrieval_attempts": 1582,
        "released_freeze_all_cycles": 1334,
        "cancelled_retrievals_absent_from_workbook": 248,
        "total_oocytes_in_valid_cohort": int(cohort.eggs.sum()),
        "total_mii_in_valid_cohort": int(cohort.mii.sum()),
        "aggregate_maturation_rate": float(cohort.mii.sum() / cohort.eggs.sum()),
        "cycles_with_at_least_10_eggs": int(cohort.eggs.ge(10).sum()),
        "cycles_with_exactly_17_eggs": int(cohort.eggs.eq(17).sum()),
        "published_retrieval_attempts_before_exclusions": 1758,
        "published_excluded_before_analysis": 176,
        "published_exclusion_breakdown": {
            "natural_cycle_ivf": 3,
            "fresh_blastocyst_transfer": 67,
            "maternal_age_over_42": 99,
            "no_sperm": 7,
        },
        "published_cancellation_breakdown": {
            "no_oocytes": 20,
            "no_mature_oocytes": 11,
            "no_fertilized_zygotes": 33,
            "embryo_developmental_arrest": 184,
        },
        "pre_trigger_removals_reported_without_count": True,
        "predictor_plausibility_flags": implausible,
    }
    if digest != "2ec6f6cd932de7e27886f99a8e81141532794a16dd259a31f6ded6257e5cf35e":
        raise RuntimeError(
            "Antalya workbook checksum does not match the frozen specification"
        )
    if len(main_first) != 1324 or len(no_fet_raw) != 10:
        raise RuntimeError(
            "Antalya worksheet linkage counts differ from the frozen audit"
        )
    if (
        cohort.retrieval_id.nunique() != 1334
        or cohort.Patient_Code.nunique() != 1315
        or couple_counts.gt(1).sum() != 19
    ):
        raise RuntimeError(
            "Antalya couple/retrieval linkage differs from the amended audit"
        )
    if len(rate_diff) and rate_diff.max() > 1e-10:
        raise RuntimeError("The reported maturation rate is not icsi/eggs")
    return (cohort, audit)


SOURCE = Path("preparation/antalya/antalya_external_clean_cohort.csv")
WORKBOOK = Path("raw/antalya/electronic_data_supplementary_material.xlsx")
SOURCE_SHA256 = "1f325beb6f5b9519d7a091056d9375e5c3349c137631842d28d75fe8b144065b"
WORKBOOK_SHA256 = "2ec6f6cd932de7e27886f99a8e81141532794a16dd259a31f6ded6257e5cf35e"
TRAIN_INDEX_SHA256 = "16dc7082360b005e0308867fcdad48d5f002a5ea5ad8415f3b313b520507ddab"
HOLDOUT_SEED = 20260904
FOLD_SEED = 202609080
UPSTREAM_NUM = [
    "age",
    "inf_years",
    "afc",
    "bmi",
    "previous_et_number",
    "stimulation_days",
    "daily_fsh",
    "daily_hmg",
    "total_fsh",
]
UPSTREAM_CAT = ["primary_infertility", "diagnosis", "stimulation_type", "trigger_type"]
COUNT_COLS = ["eggs", "mii", "fertilized", "blastocysts", "frozen_blastocysts"]
MASK_COLS = [
    "first_fet_transferred_observed",
    "first_fet_endo_observed",
    "first_fet_pregnancy_observed",
    "first_fet_live_birth_observed",
    "first_fet_eligible",
]
OUTCOME_COLS = ["first_fet_pregnancy", "first_fet_live_birth"]
NUM = (
    UPSTREAM_NUM
    + COUNT_COLS
    + ["first_fet_transferred", "first_fet_endo_mm"]
    + OUTCOME_COLS
    + MASK_COLS
)
CAT = UPSTREAM_CAT + ["first_fet_embryo_score"]
COLS = NUM + CAT
RAW_MAP = {
    "2PN": "fertilized",
    "blast>=2bb": "blastocysts",
    "blast FZ>=2bb": "frozen_blastocysts",
    "#blast transf": "first_fet_transferred",
    "endo mm": "first_fet_endo_mm",
    "PR": "first_fet_pregnancy",
    "LB 1": "first_fet_live_birth",
}
UPSTREAM_RAW = {
    "age": "age",
    "inf yrs": "inf_years",
    "afc": "afc",
    "bmi": "bmi",
    "prev ET numb": "previous_et_number",
    "stim days": "stimulation_days",
    "FSH": "daily_fsh",
    "hmg": "daily_hmg",
    "tot FSH": "total_fsh",
    "prime inf": "primary_infertility",
    "diag": "diagnosis",
    "stim type": "stimulation_type",
    "trigger type": "trigger_type",
    "eggs": "eggs",
    "ıcsı": "mii",
}


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1048576), b""):
            digest.update(block)
    return digest.hexdigest()


def _number(value, field, *, allow_missing=True):
    if value is None or (isinstance(value, str) and value.strip() in ["", "?"]):
        if allow_missing:
            return np.nan
        raise ValueError("Missing required numeric field: " + field)
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("Nonfinite source value: " + field)
    return number


def _category(value):
    return (
        "__MISSING__"
        if value is None or not str(value).strip()
        else str(value).strip().casefold()
    )


def _patient(pno):
    value = _number(pno, "pno", allow_missing=False)
    return f"ANT-{math.floor(value):05d}"


def _key(pno, date):
    return (float(pno), pd.Timestamp(date).isoformat())


def _extract_workbook(workbook, training):
    import openpyxl

    wb = openpyxl.load_workbook(workbook, read_only=False, data_only=True)
    training_ids = set(training.Patient_Code.astype(str))
    expected = {
        _key(row.source_cycle_id, row.retrieval_date): i
        for i, row in training.iterrows()
    }
    if len(expected) != len(training):
        raise ValueError("Duplicate training retrieval keys")
    collected, origins = ({}, {})
    audit = {
        "nonselected_rows_skipped_before_clinical_access": {},
        "selected_retrieval_rows": {},
        "later_training_fet_rows_skipped": 0,
        "nonselected_clinical_or_outcome_values_selected": 0,
        "holdout_values_selected": False,
    }
    for sheet_name, header_row in [
        ("tot cycles LB outcome", 2),
        ("No FET plus unknown", 1),
    ]:
        ws = wb[sheet_name]
        headers = {
            str(ws.cell(header_row, c).value): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(header_row, c).value is not None
        }
        needed = set(RAW_MAP) | set(UPSTREAM_RAW) | {"pno", "tarih", "FET#"}
        if not needed <= set(headers):
            raise ValueError("Source headers differ: " + sheet_name)
        skipped = selected = 0
        for row in range(header_row + 1, ws.max_row + 1):
            pno = ws.cell(row, headers["pno"]).value
            if pno is None:
                raise ValueError("Unexpected missing row identifier")
            if _patient(pno) not in training_ids:
                skipped += 1
                continue
            fet_number = _number(ws.cell(row, headers["FET#"]).value, "FET#")
            if sheet_name == "tot cycles LB outcome" and fet_number != 1:
                audit["later_training_fet_rows_skipped"] += 1
                continue
            date = ws.cell(row, headers["tarih"]).value
            key = _key(pno, date)
            if key not in expected or key in collected:
                raise ValueError("Unmatched or duplicate selected training retrieval")
            if fet_number != 1:
                raise ValueError("First-FET source scope changed")
            item = {}
            for raw, name in UPSTREAM_RAW.items():
                value = ws.cell(row, headers[raw]).value
                item[name] = (
                    _category(value) if name in UPSTREAM_CAT else _number(value, raw)
                )
            for raw, name in RAW_MAP.items():
                item[name] = _number(ws.cell(row, headers[raw]).value, raw)
            score_name = (
                "SET emb score" if sheet_name == "tot cycles LB outcome" else "SET emb"
            )
            score = ws.cell(row, headers[score_name]).value
            item["first_fet_embryo_score"] = _category(score)
            collected[key] = item
            origins[key] = {
                "sheet": sheet_name,
                "xlsx_row": row,
                "first_fet_number": fet_number,
            }
            selected += 1
        audit["nonselected_rows_skipped_before_clinical_access"][sheet_name] = skipped
        audit["selected_retrieval_rows"][sheet_name] = selected
    wb.close()
    if set(collected) != set(expected):
        raise ValueError("Incomplete exact training retrieval linkage")
    extended = pd.DataFrame(
        [
            collected[_key(r.source_cycle_id, r.retrieval_date)]
            for _, r in training.iterrows()
        ],
        index=training.index,
    )
    for c in UPSTREAM_NUM + ["eggs", "mii"]:
        if not np.allclose(
            extended[c], training[c], rtol=0, atol=1e-12, equal_nan=True
        ):
            raise ValueError("Original workbook versus development data mismatch: " + c)
    for c in UPSTREAM_CAT:
        if (
            extended[c].tolist()
            != training[c].fillna("__MISSING__").astype(str).tolist()
        ):
            raise ValueError("Original category mapping differs: " + c)
    audit["linkage_rule"] = (
        "exact numeric pno plus retrieval date; patient ANT-floor(pno), fixed Cycle_ID"
    )
    audit["linkage_sha256"] = hashlib.sha256(
        json.dumps(
            [
                {"Cycle_ID": training.loc[expected[k], "Cycle_ID"], **origins[k]}
                for k in sorted(expected)
            ],
            sort_keys=True,
        ).encode()
    ).hexdigest()
    audit["upstream_parity_atol"] = 1e-12
    return (extended, audit)


def _load_selected_pool(source=SOURCE, workbook=WORKBOOK):
    if sha(source) != SOURCE_SHA256 or sha(workbook) != WORKBOOK_SHA256:
        raise ValueError("Pinned original Antalya inputs changed")
    identifiers = pd.read_csv(source, usecols=["Patient_Code"])
    if (
        len(identifiers) != 1334
        or identifiers.Patient_Code.nunique() != 1315
        or identifiers.Patient_Code.isna().any()
    ):
        raise ValueError("Source identifier inventory changed")
    split = GroupShuffleSplit(n_splits=1, test_size=0.2, random_state=HOLDOUT_SEED)
    ti, hi = next(split.split(identifiers, groups=identifiers.Patient_Code))
    keep = sorted(map(int, ti))
    train_ids = set(identifiers.iloc[ti].Patient_Code)
    if (
        len(ti),
        len(train_ids),
        len(hi),
        identifiers.iloc[hi].Patient_Code.nunique(),
    ) != (1067, 1052, 267, 263):
        raise ValueError("Development split cardinality changed")
    index_hash = hashlib.sha256(np.asarray(keep, dtype="<i8").tobytes()).hexdigest()
    if index_hash != TRAIN_INDEX_SHA256:
        raise ValueError("Exact development indices changed")
    keep_set = set(keep)
    source_columns = UPSTREAM_NUM + UPSTREAM_CAT + ["eggs", "mii"]
    frame = pd.read_csv(
        source,
        usecols=["Patient_Code", "source_cycle_id", "retrieval_date"] + source_columns,
        skiprows=lambda row: row > 0 and row - 1 not in keep_set,
    )
    if frame.Patient_Code.tolist() != identifiers.iloc[keep].Patient_Code.tolist():
        raise ValueError("Development row order changed")
    for c in UPSTREAM_NUM + ["eggs", "mii"]:
        frame[c] = pd.to_numeric(frame[c], errors="raise")
    for c in UPSTREAM_CAT:
        frame[c] = frame[c].fillna("__MISSING__").astype(str)
    frame["Patient_Code"] = frame.Patient_Code.astype(str)
    frame["Cycle_ID"] = ["A" + str(i) for i in keep]
    ext, workbook_audit = _extract_workbook(workbook, frame)
    for c in COUNT_COLS[2:] + [
        "first_fet_transferred",
        "first_fet_endo_mm",
        "first_fet_pregnancy",
        "first_fet_live_birth",
        "first_fet_embryo_score",
    ]:
        frame[c] = ext[c]
    for c in COUNT_COLS:
        if not _valid_count(frame[c]).all():
            raise ValueError("Missing or invalid source count: " + c)
        frame[c] = frame[c].astype(int)
    if not frame.eggs.gt(0).all() or not frame.mii.le(frame.eggs).all():
        raise ValueError("Original retrieval/MII domain changed")
    for value, mask in [
        ("first_fet_transferred", "first_fet_transferred_observed"),
        ("first_fet_endo_mm", "first_fet_endo_observed"),
        ("first_fet_pregnancy", "first_fet_pregnancy_observed"),
        ("first_fet_live_birth", "first_fet_live_birth_observed"),
    ]:
        present = frame[value].notna()
        if value in OUTCOME_COLS and (not frame.loc[present, value].isin([0, 1]).all()):
            raise ValueError("Unknown binary outcome code: " + value)
        if value == "first_fet_transferred" and (
            not frame.loc[present, value].isin([0, 1, 2]).all()
        ):
            raise ValueError("Unknown transfer count")
        frame[mask] = present.astype(int)
        frame[value] = frame[value].fillna(0.0)
    frame["first_fet_eligible"] = (
        frame.first_fet_transferred_observed.eq(1)
        & frame.first_fet_transferred.isin([1, 2])
    ).astype(int)
    frame = frame[["Patient_Code", "Cycle_ID"] + COLS].copy()
    receipt = dict(
        holdout_values_selected=False,
        source_files_unchanged=sha(source) == SOURCE_SHA256
        and sha(workbook) == WORKBOOK_SHA256,
    )
    if not receipt["source_files_unchanged"]:
        raise ValueError("Original input bytes changed during audit")
    expected_size = (1067, 1052)
    if (len(frame), frame.Patient_Code.nunique()) != expected_size:
        raise ValueError("Selected source cohort cardinality changed")
    return (frame, receipt)


def _valid_count(values):
    x = pd.to_numeric(values, errors="coerce")
    return x.notna() & np.isfinite(x) & x.ge(0) & x.eq(np.rint(x))


def repair(frame, train):
    if not set(COLS) <= set(frame) or not set(COLS) <= set(train):
        raise ValueError("Incomplete shared synthetic schema")
    out = frame.copy()
    diag = {
        "numeric_imputed": {},
        "numeric_clipped": {},
        "categorical_repaired": {},
        "raw_invalid_count_fraction": {},
        "repair_scope": "training marginal support; eggs>=1/MII<=eggs; masks and explicit dummy values; recorded upstream-count exceptions retained",
    }
    for c in COUNT_COLS:
        diag["raw_invalid_count_fraction"][c] = float((~_valid_count(out[c])).mean())
    diag["raw_mii_gt_eggs_fraction"] = float(
        pd.to_numeric(out.mii, errors="coerce")
        .gt(pd.to_numeric(out.eggs, errors="coerce"))
        .mean()
    )
    for c in NUM:
        reference = pd.to_numeric(train[c], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )
        if not reference.notna().any():
            raise ValueError("Training numeric field has no observed support: " + c)
        value = pd.to_numeric(out[c], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )
        lo, hi = (
            (0.0, 1.0)
            if c in MASK_COLS
            else (float(reference.min()), float(reference.max()))
        )
        diag["numeric_imputed"][c] = int(value.isna().sum())
        diag["numeric_clipped"][c] = int((value.lt(lo) | value.gt(hi)).sum())
        fill = 0.0 if c in MASK_COLS else float(reference.median())
        out[c] = value.fillna(fill).clip(lo, hi)
    for c in COUNT_COLS:
        out[c] = out[c].round().clip(lower=0).astype(int)
    out["eggs"] = out.eggs.clip(lower=1)
    out["mii"] = np.minimum(out.mii, out.eggs).astype(int)
    for c in MASK_COLS:
        out[c] = out[c].ge(0.5).astype(int)
    out["first_fet_transferred"] = (
        out.first_fet_transferred.round().clip(0, 2).astype(int)
    )
    for value, mask in [
        ("first_fet_transferred", "first_fet_transferred_observed"),
        ("first_fet_endo_mm", "first_fet_endo_observed"),
        ("first_fet_pregnancy", "first_fet_pregnancy_observed"),
        ("first_fet_live_birth", "first_fet_live_birth_observed"),
    ]:
        raw_missing = ~np.isfinite(
            pd.to_numeric(frame[value], errors="coerce").to_numpy(float)
        )
        out.loc[raw_missing, mask] = 0
        if value in OUTCOME_COLS:
            out[value] = out[value].ge(0.5).astype(int)
        out.loc[out[mask].eq(0), value] = 0
    out["first_fet_eligible"] = (
        out.first_fet_transferred_observed.eq(1)
        & out.first_fet_transferred.isin([1, 2])
    ).astype(int)
    for c in CAT:
        ref = train[c].fillna("__MISSING__").astype(str)
        v = out[c].fillna("__MISSING__").astype(str)
        known = v.isin(set(ref))
        diag["categorical_repaired"][c] = int((~known).sum())
        out[c] = v.where(known, ref.mode().iloc[0])
    return (out, diag)


def metadata_for(data):
    from sdv.metadata import Metadata

    meta = Metadata.detect_from_dataframes(data, infer_sdtypes=True, infer_keys=None)
    for table, frame in data.items():
        for col in frame:
            typ = (
                "id"
                if col in ["Patient_Code", "Cycle_ID"]
                else "categorical"
                if col in CAT
                else "numerical"
            )
            meta.update_column(col, table_name=table, sdtype=typ)
        if table == "patients":
            meta.set_primary_key("Patient_Code", table_name=table)
        elif table == "retrievals":
            meta.set_primary_key("Cycle_ID", table_name=table)
    if "patients" in data:
        meta.add_relationship(
            parent_table_name="patients",
            child_table_name="retrievals",
            parent_primary_key="Patient_Code",
            child_foreign_key="Patient_Code",
        )
    meta.validate()
    meta.validate_data(data)
    return meta
